"""Independent synthetic DOCX extraction, old-source red/green and real masking chain."""
import contextlib,gc,io,json,logging,pathlib,subprocess,sys,tempfile,types
ROOT=pathlib.Path(__file__).resolve().parents[4];sys.path[:0]=[str(ROOT/'backend'),str(ROOT/'audyt')]
assert subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()=='57e39fbb6e30cccdce87103b9f72f066c4126c2a'
logging.disable(logging.CRITICAL)
import document_processor as dp
from docx import Document
from pipeline_new import run_pipeline_express
from pipeline import pseudonymize_document
from unmask_sim import unmask
import engine_quality as eq
old=types.ModuleType('base_document_processor');old.__file__=str(ROOT/'backend/document_processor.py')
src=subprocess.check_output(['git','show','f95d1b9:backend/document_processor.py'],text=True,encoding='utf-8')
with contextlib.redirect_stdout(io.StringIO()):exec(compile(src,old.__file__,'exec'),old.__dict__)
def data(d):
    b=io.BytesIO();d.save(b);return b.getvalue()
# Order/empty/multiple paragraphs/horizontal+vertical merging, with tables before/between paragraphs.
d=Document();t=d.add_table(rows=1,cols=2);t.cell(0,0).text='start_left';t.cell(0,1).text='start_right'
d.add_paragraph('before');t=d.add_table(rows=3,cols=3)
t.cell(0,0).merge(t.cell(0,1)).text='horizontal';t.cell(0,2).text='right_top'
t.cell(1,0).merge(t.cell(2,0)).text='vertical';t.cell(1,1).text='p1';t.cell(1,1).add_paragraph('p2')
t.cell(2,1).text='left_bottom';t.cell(2,2).text='right_bottom'
d.add_paragraph('between');d.add_table(rows=1,cols=1).cell(0,0).text='end_table';d.add_paragraph('after')
blob=data(d);expected='\n'.join(['start_left','start_right','before','horizontal','right_top','vertical','p1','p2','left_bottom','right_bottom','between','end_table','after'])
assert dp._extract_text(blob,'synthetic.docx')==(expected,None,None)
assert old._extract_text(blob,'synthetic.docx')==('before\nbetween\nafter',None,None)
print('PASS structural ordered body matrix; exact old-source RED -> delivered GREEN')
# Rectangular merge across rows+columns, no duplicate aliases.
d=Document();t=d.add_table(rows=3,cols=3);t.cell(0,0).merge(t.cell(1,1)).text='block'
for (r,c),v in [((0,2),'right0'),((1,2),'right1'),((2,0),'left2'),((2,1),'mid2'),((2,2),'right2')]:t.cell(r,c).text=v
assert dp._extract_text(data(d),'synthetic.docx')==('block\nright0\nright1\nleft2\nmid2\nright2',None,None)
print('PASS rectangular merge once, body-only table accepted')
assert dp._extract_text(b'not a valid archive','synthetic.docx')[0]=='' and dp._extract_text(b'not a valid archive','synthetic.docx')[1]
assert dp._extract_text(b'plain text','synthetic.txt')==old._extract_text(b'plain text','synthetic.txt')
from reportlab.pdfgen import canvas
pdf=io.BytesIO();c=canvas.Canvas(pdf);c.drawString(72,700,'synthetic PDF control');c.save()
assert dp._extract_text(pdf.getvalue(),'synthetic.pdf')==old._extract_text(pdf.getvalue(),'synthetic.pdf')
print('PASS corrupt DOCX error and unchanged TXT/digital PDF controls')
with tempfile.TemporaryDirectory(prefix='lynx_docx_independent_') as tmp:
    with contextlib.redirect_stdout(io.StringIO()):state=eq.build_state(tmp)
    emails=['before.audit@example.pl','cell.one@example.pl','cell.two@example.pl','after.audit@example.pl']
    d=Document();d.add_paragraph('Kontakt: '+emails[0]);t=d.add_table(rows=2,cols=2)
    t.cell(0,0).text=emails[1];t.cell(0,1).text='526-000-12-46';t.cell(1,0).text=emails[2]
    d.add_paragraph('Kontakt: '+emails[3]);blob=data(d)
    text,error,meta=dp._extract_text(blob,'synthetic.docx');assert error is None and meta is None
    for value in emails+['526-000-12-46']:assert text.count(value)==1
    assert text.index(emails[0])<text.index(emails[1])<text.index(emails[2])<text.index(emails[3])
    with contextlib.redirect_stdout(io.StringIO()):
        result=pseudonymize_document(text,state)
        out,rev,blocked=run_pipeline_express(text,{},state.anonymizer)
    assert not result.error and not result.guard_blocked and not blocked
    expected_values=set(emails+['526-000-12-46'])
    for masked,mapping in [(result.text,result.reverse_map),(out,rev)]:
        assert set(mapping.values())==expected_values and len(mapping)==5
        assert not any(value in masked for value in expected_values)
        assert unmask(masked,mapping)==text
    print('PASS synthetic DOCX 2x2 -> extraction -> full production wrapper/Guard and Express; complete maps, restoration 1:1')
    state=None;gc.collect()