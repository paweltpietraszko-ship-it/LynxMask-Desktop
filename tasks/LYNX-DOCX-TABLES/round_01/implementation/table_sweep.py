# LYNX-DOCX-TABLES: pary wartości w sąsiednich komórkach (ten sam wiersz) i wierszach (ta sama kolumna) tabeli DOCX,
# produkcyjny _extract_text + pipeline, oba tryby, NER wyłączony, dane syntetyczne. Uruchom z backend/.
import contextlib, io, itertools, re, sys
from unittest.mock import patch
sys.path[:0] = [".", "../audyt"]
V = {"EMAIL1": "k.zalewska@example.pl", "EMAIL2": "m.nowicka@example.pl", "TEL": "512 987 654", "NIP": "526-000-12-46",
     "PESEL": "44051401359", "IBAN": "PL61 1090 1014 0000 0712 1981 2874", "KWOTA": "1 250,00 zł", "TXT": "Opis pozycji"}

def docx_bytes(a, b, horizontal):
    import docx
    d = docx.Document()
    d.add_paragraph("Zestawienie:")
    t = d.add_table(rows=1 if horizontal else 2, cols=2 if horizontal else 1)
    (t.rows[0].cells[0] if horizontal else t.rows[0].cells[0]).text = a
    (t.rows[0].cells[1] if horizontal else t.rows[1].cells[0]).text = b
    d.add_paragraph("Koniec zestawienia.")
    buf = io.BytesIO(); d.save(buf); return buf.getvalue()

def run():
    import layers.ner_adapter as na
    from document_processor import _extract_text
    from pipeline_new import run_pipeline_new, run_pipeline_express
    from unmask_sim import unmask
    bad, n = [], 0
    for (ka, kb), horiz in itertools.product(itertools.permutations(V, 2), (True, False)):
        text, err, _ = _extract_text(docx_bytes(V[ka], V[kb], horiz), "t.docx")
        for fn in (run_pipeline_new, run_pipeline_express):
            n += 1
            with patch.object(na, "process_ner", return_value=({}, {})), contextlib.redirect_stdout(io.StringIO()):
                out, rev, _ = fn(text, anon_map={})
            frags = [w for k in (ka, kb) if k != "TXT" for w in re.findall(r"[A-Za-z0-9]{3,}", V[k]) if w != "example"]
            problem = (err or V[ka] not in text or V[kb] not in text or any(f in out for f in frags) or "_NUMER" in out
                       or out.count("\n") != text.count("\n") or "Koniec zestawienia." not in unmask(out, rev))
            if problem:
                bad.append((ka, kb, "poziomo" if horiz else "pionowo", fn.__name__, out))
    return n, bad

if __name__ == "__main__":
    n, bad = run()
    print(f"przypadków {n}: problemy {len(bad)}")
    for b in bad[:12]:
        print("  ", b)
