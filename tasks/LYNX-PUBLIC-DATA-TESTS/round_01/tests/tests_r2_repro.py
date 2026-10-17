"""Independent extraction-occurrence check; synthetic fixtures and real production chain."""
import contextlib,gc,io,json,logging,pathlib,subprocess,sys,tempfile
ROOT=pathlib.Path(__file__).resolve().parents[4]
sys.path[:0]=[str(ROOT/'backend'),str(ROOT/'audyt')]
assert subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()=='d7a08689af663f422e8ed7e3cb1d5a204e3b46f3'
import engine_quality as eq
from document_processor import _extract_text
from docx import Document
import pipeline
logging.disable(logging.CRITICAL)
ok,why=eq._engine_available()
if not ok: raise SystemExit('NOT_RUN: '+why)
with tempfile.TemporaryDirectory(prefix='lynx_layout_independent_') as temp:
    root=pathlib.Path(temp)
    st=eq.build_state(str(root/'profile'))
    def run(text):
        with contextlib.redirect_stdout(io.StringIO()): return pipeline.pseudonymize_document(text,st)
    value='audyt.layout@example.pl'
    cases=[]
    for repeat in [False,True]:
        d=Document(); d.add_paragraph('Kontakt: '+value)
        cell=value if repeat else 'inna.komorka@example.pl'
        d.add_table(rows=1,cols=1).cell(0,0).text=cell
        path=root/('repeated.docx' if repeat else 'unique.docx');d.save(path)
        extracted,error,_=_extract_text(path.read_bytes(),path.name)
        assert error is None and len(d.tables)==1 and d.tables[0].cell(0,0).text==cell
        assert extracted=='Kontakt: '+value and extracted.count(value)==1
        entry={'file':path.name,'type':'docx','fields':[{'category':'EMAIL','value':value},{'category':'EMAIL','value':cell,'in_table':True}]}
        (root/'manifest.json').write_text(json.dumps([entry]),encoding='utf-8')
        rep=eq.measure_layout(str(root),run)
        assert rep['by_type_location']['docx/tekst']=={'masked':1}
        table=rep['by_type_location']['docx/tabela']
        if repeat:
            assert table=={'masked':1},rep
            print('REPRO: missing table occurrence is falsely masked when paragraph contains the same value')
        else:
            assert table=={'EXTRACTION_MISSING':1},rep
            print('CONTROL: unique missing table value correctly EXTRACTION_MISSING')
        dump=json.dumps(rep)
        assert value not in dump and cell not in dump and path.name not in dump
        cases.append({'case':'repeat' if repeat else 'unique','status':rep['status'],'by_type_location':rep['by_type_location'],'m3':rep['docs'][0]['m3']})
    print(json.dumps(cases))
    st=None;gc.collect()