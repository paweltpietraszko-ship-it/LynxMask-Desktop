"""Narrow independent F1 recheck; repeat and overlapping-value siblings, synthetic only."""
import contextlib,gc,io,json,logging,pathlib,subprocess,sys,tempfile
ROOT=pathlib.Path(__file__).resolve().parents[4]
sys.path[:0]=[str(ROOT/'backend'),str(ROOT/'audyt')]
assert subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()=='638b31813962d2ba071ac02bc7d21cb98481dd9a'
import engine_quality as eq
from docx import Document
from document_processor import _extract_text
import pipeline
logging.disable(logging.CRITICAL)
ok,why=eq._engine_available()
if not ok: raise SystemExit('NOT_RUN: '+why)
with tempfile.TemporaryDirectory(prefix='lynx_layout_recheck_') as temp:
    root=pathlib.Path(temp); st=eq.build_state(str(root/'profile'))
    def run(text):
        with contextlib.redirect_stdout(io.StringIO()): return pipeline.pseudonymize_document(text,st)
    short='audyt.layout@example.pl';long='biuro.'+short
    cases=[('original-repeat',[short],short,{'EXTRACTION_UNVERIFIED':1}),
           ('unique-control',[short],'inna.komorka@example.pl',{'EXTRACTION_MISSING':1}),
           ('both-present',[short,short],None,None),
           ('overlap-sibling',[long],short,None)]
    for label,paragraphs,cell,expected in cases:
        d=Document()
        fields=[]
        for v in paragraphs:
            d.add_paragraph('Kontakt: '+v)
            fields.append({'category':'EMAIL','value':v})
        if cell is not None:
            d.add_table(rows=1,cols=1).cell(0,0).text=cell
            fields.append({'category':'EMAIL','value':cell,'in_table':True})
        path=root/'sample.docx';d.save(path)
        text,error,_=_extract_text(path.read_bytes(),path.name)
        assert error is None
        assert text=='\n'.join('Kontakt: '+v for v in paragraphs)
        (root/'manifest.json').write_text(json.dumps([{'file':path.name,'type':'docx','fields':fields}]),encoding='utf-8')
        rep=eq.measure_layout(str(root),run)
        by=rep['by_type_location']
        if expected is not None:
            assert by['docx/tabela']==expected,rep
        if label=='original-repeat': assert by['docx/tekst']=={'EXTRACTION_UNVERIFIED':1}
        if label=='unique-control': assert by['docx/tekst']=={'masked':1}
        if label=='both-present': assert by=={'docx/tekst':{'masked':2}}
        if label=='overlap-sibling':
            assert len(d.tables)==1 and d.tables[0].cell(0,0).text==short
            assert text.count(short)==1 and short!=long
            assert by['docx/tabela']=={'masked':1},rep
            print('REPRO residual F1: unread cell value occurs only inside a different, longer paragraph email; table falsely masked')
        assert not any(v in json.dumps(rep) for v in [short,long,path.name])
        print(label,json.dumps(by),'M3',rep['docs'][0]['m3'])
    st=None;gc.collect()